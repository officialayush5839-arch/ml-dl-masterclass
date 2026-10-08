import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Reference exact palette
    BG_COLOR = RGBColor(7, 9, 18)       # #070912
    CARD_BG = RGBColor(13, 17, 30)      # #0d111e
    ORANGE = RGBColor(255, 85, 0)       # #ff5500
    LIME = RGBColor(163, 230, 53)       # #a3e635
    WHITE = RGBColor(248, 250, 252)     # #f8fafc
    SLATE = RGBColor(203, 213, 225)     # #cbd5e1
    SLATE_MUTED = RGBColor(148, 163, 184)# #94a3b8

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title, subtitle=None):
        add_bg(slide)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title.upper()
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = WHITE

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = ORANGE
            p_sub.space_before = Pt(4)

    def add_lime_card(slide, left, top, width, height, title, subtitle):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = LIME
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE

        if subtitle:
            p_s = tf.add_paragraph()
            p_s.text = subtitle
            p_s.font.size = Pt(12)
            p_s.font.color.rgb = SLATE
            p_s.space_before = Pt(4)

    # -----------------------------------------------------------------
    # SLIDE 1: MACHINE LEARNING (Title)
    # -----------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(7.5), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "MACHINE\nLEARNING"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf1.add_paragraph()
    p2.text = "An Introduction to ML, Deep Learning & the AI Ecosystem\nTeaching Machines to Learn from Data"
    p2.font.size = Pt(14)
    p2.font.color.rgb = SLATE_MUTED
    p2.space_before = Pt(16)

    p3 = tf1.add_paragraph()
    p3.text = "Presented by Ayush Vishwakarma"
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = ORANGE
    p3.space_before = Pt(14)

    # Bottom bar steps
    steps = ["01 DATA", "02 LEARNING", "03 MODEL", "04 PREDICTION", "05 DECISION"]
    for i, step in enumerate(steps):
        s_box = s1.shapes.add_textbox(Inches(0.8 + i * 2.3), Inches(6.2), Inches(2.1), Inches(0.6))
        stf = s_box.text_frame
        sp = stf.paragraphs[0]
        sp.text = step
        sp.font.size = Pt(12)
        sp.font.bold = True
        sp.font.color.rgb = ORANGE

    # -----------------------------------------------------------------
    # SLIDE 2: Can a machine learn without being explicitly programmed?
    # -----------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Can a machine learn without\nbeing explicitly programmed?")
    col_w = Inches(3.6)
    gap = Inches(0.4)
    add_lime_card(s2, Inches(0.8), Inches(2.2), col_w, Inches(1.6), "YouTube / Netflix", "What should you watch next?")
    add_lime_card(s2, Inches(0.8 + (col_w+gap)), Inches(2.2), col_w, Inches(1.6), "Email", "Is this spam?")
    add_lime_card(s2, Inches(0.8 + (col_w+gap)*2), Inches(2.2), col_w, Inches(1.6), "Phone Camera", "What is in this image?")

    q_box = s2.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.5))
    q_tf = q_box.text_frame
    qp = q_tf.paragraphs[0]
    qp.text = "How does the machine know?\n\nA. Someone programmed every possibility\nB. The machine learned patterns from data\nC. The machine guesses randomly\n\nReveal: B · It learns patterns from data."
    qp.font.size = Pt(14)
    qp.font.color.rgb = WHITE

    # -----------------------------------------------------------------
    # SLIDE 3: What actually is Machine Learning?
    # -----------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "What actually is Machine Learning?")
    c_w = Inches(5.6)
    add_lime_card(s3, Inches(0.8), Inches(2.0), c_w, Inches(2.4), "Traditional Programming", "RULES + DATA\n↓\nOUTPUT\n\nIf temperature > 30°C → Turn Fan ON")
    add_lime_card(s3, Inches(6.8), Inches(2.0), c_w, Inches(2.4), "Machine Learning", "DATA + EXAMPLES\n↓\nMODEL → PREDICTION\n\nHouse Features → TrainingData → ML Model → Predicted Price")

    d_box = s3.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11.7), Inches(2.0))
    d_tf = d_box.text_frame
    dp = d_tf.paragraphs[0]
    dp.text = "ML = Learning patterns from data to make predictions or decisions.\n\n10,000 pictures of cats and dogs. Rules for every possible cat?\nAnswer: No. That is where ML becomes useful."
    dp.font.size = Pt(14)
    dp.font.color.rgb = WHITE

    # -----------------------------------------------------------------
    # SLIDE 4: HOW DOES ML LEARN?
    # -----------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "HOW DOES ML LEARN?")
    c4_w = Inches(3.6)
    c4_gap = Inches(0.4)
    add_lime_card(s4, Inches(0.8), Inches(2.0), c4_w, Inches(3.6), "SUPERVISED", "Learn from labeled examples.\n\nInput → Correct Answer\n\n• Spam / Not Spam\n• House-price prediction\n• Image classification\n• Disease prediction")
    add_lime_card(s4, Inches(0.8 + (c4_w+c4_gap)), Inches(2.0), c4_w, Inches(3.6), "UNSUPERVISED", "Find patterns without labels.\n\nInput → Discover Structure\n\n• Customer segmentation\n• Anomaly detection\n• Grouping documents\n• Recommendation systems")
    add_lime_card(s4, Inches(0.8 + (c4_w+c4_gap)*2), Inches(2.0), c4_w, Inches(3.6), "REINFORCEMENT", "Learn through rewards & penalties.\n\nAction → Reward / Penalty → Better Action\n\n• Robotics\n• Game agents\n• Autonomous systems\n• Optimization")

    q4_box = s4.shapes.add_textbox(Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.0))
    q4_tf = q4_box.text_frame
    q4p = q4_tf.paragraphs[0]
    q4p.text = "Predict whether a student will pass using attendance, marks and study hours. Which type?\nAnswer: SUPERVISED LEARNING"
    q4p.font.size = Pt(13)
    q4p.font.bold = True
    q4p.font.color.rgb = ORANGE

    # -----------------------------------------------------------------
    # SLIDE 5: THE AI FAMILY TREE
    # -----------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "THE AI FAMILY TREE")
    t5_box = s5.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.5))
    t5_tf = t5_box.text_frame
    t5p = t5_tf.paragraphs[0]
    t5p.text = "Artificial Intelligence\n└── Machine Learning\n    ├── Traditional ML\n    └── Deep Learning\n        ├── CNNs\n        ├── RNNs\n        └── Transformers\n            ├── NLP & LLMs\n            └── Generative AI"
    t5p.font.size = Pt(16)
    t5p.font.color.rgb = SLATE

    t5_right = s5.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(5.6), Inches(4.5))
    t5_rtf = t5_right.text_frame
    rp = t5_rtf.paragraphs[0]
    rp.text = "Deep Learning\nNeural networks with many layers learn complex representations from large amounts of data.\n\nGenerative AI\nModels that generate new content instead of only classifying or predicting.\n\nTEXT · IMAGES · AUDIO · VIDEO · CODE\n\nChatGPT writes an essay. Is that classification?\nAnswer: No — it is generative modeling."
    rp.font.size = Pt(14)
    rp.font.color.rgb = WHITE

    # -----------------------------------------------------------------
    # SLIDE 6: MEET GANs
    # -----------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "MEET GANs", "GENERATIVE ADVERSARIAL NETWORKS")
    add_lime_card(s6, Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.2), "GENERATOR ➔ DISCRIMINATOR", "Creates fake samples ➔ Real or Fake?\n\nCOMPETE ➔ LEARN ➔ IMPROVE\n'Here is a fake painting.' → 'That is fake.' → 'I will make a better one.'")
    add_lime_card(s6, Inches(6.8), Inches(2.0), Inches(5.6), Inches(2.2), "Which image was generated?", "Sample A vs Sample B\n\nGuess before the reveal.\nReveal: both samples are AI-generated.")

    g_note = s6.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11.7), Inches(2.0))
    g_tf = g_note.text_frame
    gp = g_tf.paragraphs[0]
    gp.text = "Eventually: highly realistic outputs (Images · Audio · Video · Text · Code)\n\nGANs are an important generative-model family; modern generative AI also includes diffusion and transformer-based models."
    gp.font.size = Pt(13)
    gp.font.color.rgb = SLATE

    # -----------------------------------------------------------------
    # SLIDE 7: ML IS NOT ONE INDUSTRY
    # -----------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "ML IS NOT ONE INDUSTRY.", "IT'S A LAYER ACROSS INDUSTRIES.")
    ind_w = Inches(3.6)
    ind_gap = Inches(0.4)
    add_lime_card(s7, Inches(0.8), Inches(2.0), ind_w, Inches(1.8), "HEALTHCARE", "Diagnosis · Drug discovery · Medical imaging")
    add_lime_card(s7, Inches(0.8 + (ind_w+ind_gap)), Inches(2.0), ind_w, Inches(1.8), "FINANCE", "Fraud detection · Risk scoring · Algorithmic systems")
    add_lime_card(s7, Inches(0.8 + (ind_w+ind_gap)*2), Inches(2.0), ind_w, Inches(1.8), "AUTOMOTIVE", "ADAS · Autonomous driving · Predictive maintenance")

    add_lime_card(s7, Inches(0.8), Inches(4.1), ind_w, Inches(1.8), "E-COMMERCE", "Recommendations · Search · Demand forecasting")
    add_lime_card(s7, Inches(0.8 + (ind_w+ind_gap)), Inches(4.1), ind_w, Inches(1.8), "CYBERSECURITY", "Threat detection · Anomaly detection · Malware analysis")
    add_lime_card(s7, Inches(0.8 + (ind_w+ind_gap)*2), Inches(4.1), ind_w, Inches(1.8), "MEDIA", "Recommendations · Content generation · Personalization")

    c7_box = s7.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8))
    c7_p = c7_box.text_frame.paragraphs[0]
    c7_p.text = "Challenge: 30 seconds. Pick one industry. What is ONE problem ML could solve?"
    c7_p.font.size = Pt(13)
    c7_p.font.bold = True
    c7_p.font.color.rgb = ORANGE

    # -----------------------------------------------------------------
    # SLIDE 8: DON'T START WITH "I WANT TO USE AI."
    # -----------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "DON'T START WITH 'I WANT TO USE AI.'", "START WITH: 'WHAT PROBLEM DO I HAVE?'")
    pipe_steps = ["01 PROBLEM", "02 DATA", "03 CLEAN & PREP", "04 CHOOSE MODEL", "05 TRAIN", "06 EVALUATE", "07 DEPLOY", "08 MONITOR"]
    for i, st in enumerate(pipe_steps):
        b = s8.shapes.add_textbox(Inches(0.8 + i * 1.45), Inches(2.0), Inches(1.4), Inches(0.8))
        bp = b.text_frame.paragraphs[0]
        bp.text = st
        bp.font.size = Pt(10)
        bp.font.bold = True
        bp.font.color.rgb = ORANGE

    add_lime_card(s8, Inches(0.8), Inches(3.0), Inches(5.6), Inches(3.2), "Problem ➔ Approach (Part 1)", "• Predict house prices ➔ Regression\n• Detect spam ➔ Classification\n• Group customers ➔ Clustering\n• Detect unusual transactions ➔ Anomaly Detection")
    add_lime_card(s8, Inches(6.8), Inches(3.0), Inches(5.6), Inches(3.2), "Problem ➔ Approach (Part 2)", "• Recommend products ➔ Recommender System\n• Detect objects ➔ Deep Learning\n• Crop disease detection ➔ Computer Vision\n• Chatbot / document assistant ➔ NLP + LLM")

    # -----------------------------------------------------------------
    # SLIDE 9: WHAT SKILLS DO YOU ACTUALLY NEED?
    # -----------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "WHAT SKILLS DO YOU ACTUALLY NEED?")
    sk_w = Inches(2.2)
    sk_gap = Inches(0.18)
    add_lime_card(s9, Inches(0.8), Inches(2.0), sk_w, Inches(2.2), "Level 1", "Python · Git · Problem Solving")
    add_lime_card(s9, Inches(0.8 + (sk_w+sk_gap)), Inches(2.0), sk_w, Inches(2.2), "Level 2 Data", "NumPy · Pandas · SQL · Visualization")
    add_lime_card(s9, Inches(0.8 + (sk_w+sk_gap)*2), Inches(2.0), sk_w, Inches(2.2), "Level 3 ML", "Scikit-learn · Regression · Classification · Clustering")
    add_lime_card(s9, Inches(0.8 + (sk_w+sk_gap)*3), Inches(2.0), sk_w, Inches(2.2), "Level 4 DL", "PyTorch / TensorFlow · Neural Networks · CNNs · Transformers")
    add_lime_card(s9, Inches(0.8 + (sk_w+sk_gap)*4), Inches(2.0), sk_w, Inches(2.2), "Level 5 Prod", "APIs · Docker · Cloud · MLOps · Monitoring")

    sk_foot = s9.shapes.add_textbox(Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.2))
    sk_tf = sk_foot.text_frame
    skp = sk_tf.paragraphs[0]
    skp.text = "Careers: ML Engineer · Data Scientist · AI / Deep Learning Engineer · MLOps Engineer · Computer Vision Engineer · NLP / LLM Engineer · Research Engineer · AI Researcher\n\nYou do NOT need to learn everything before building your first project.\nPython → Data → Basic ML → Project → Deep Learning\n\nPoll: Which role sounds most interesting? BUILD · ANALYZE · RESEARCH · DEPLOY · CREATE"
    skp.font.size = Pt(13)
    skp.font.color.rgb = WHITE

    # -----------------------------------------------------------------
    # SLIDE 10: DON'T JUST WATCH. BUILD.
    # -----------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "DON'T JUST WATCH. BUILD.")
    add_lime_card(s10, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.5), "AUTHENTIC RESOURCES", "• Google Machine Learning Crash Course\n• Microsoft ML for Beginners\n• DeepLearning.AI Machine Learning Specialization\n• fast.ai · Practical Deep Learning for Coders\n• 3Blue1Brown · Neural Networks\n\n01 Start → 02 Learn ML → 03 Build 3–5 Projects → 04 Deep Learning → 05 Go Professional")
    add_lime_card(s10, Inches(6.8), Inches(2.0), Inches(5.6), Inches(4.5), "YOUR TURN.", "1. What problem around you generates data?\n2. Could ML learn something from that data?\n3. What would you build?\n\n'You don't need to become an AI expert to start. You need a problem, some data, and the curiosity to experiment.'\n\nAYUSH VISHWAKARMA · Machine Learning · From Data to Intelligence")

    output_path = "Machine_Learning_and_Deep_Learning_Masterclass.pptx"
    prs.save(output_path)
    print(f"Presentation regenerated: {output_path}")

if __name__ == "__main__":
    create_deck()
